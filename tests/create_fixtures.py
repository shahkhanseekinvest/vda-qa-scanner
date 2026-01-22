#!/usr/bin/env python3
"""
Create test fixtures for smoke testing.

This script generates sample files in various formats with known content
that should trigger the scanner. These fixtures verify that:
1. Each file type parser works correctly
2. Exact matching works
3. Priority prefix matching works
4. Normalized matching works (for terms with special characters)
"""

import os
from pathlib import Path

import pandas as pd
from docx import Document
from pptx import Presentation
from pptx.util import Inches


def create_fixtures_dir():
    """Create the fixtures directory structure."""
    fixtures_dir = Path(__file__).parent / "fixtures"
    vdr_dir = fixtures_dir / "sample_vdr"
    vdr_dir.mkdir(parents=True, exist_ok=True)
    return fixtures_dir, vdr_dir


def create_tracker_file(fixtures_dir):
    """Create a sample anonymization tracker Excel file."""
    # Create tracker with 'Before' column containing terms to search for
    tracker_data = {
        "Before": [
            "Acme Corporation",      # Exact match test
            "John Smith",            # Exact match test
            "(555) 123-4567",        # Normalized match test (phone number)
            "secret@email.com",      # Normalized match test (email-like)
            "123-45-6789",           # Normalized match test (SSN-like)
            "Confidential",          # Exact match test
        ],
        "After": [
            "[COMPANY]",
            "[NAME]",
            "[PHONE]",
            "[EMAIL]",
            "[SSN]",
            "[REDACTED]",
        ]
    }
    df = pd.DataFrame(tracker_data)
    
    tracker_path = fixtures_dir / "test_tracker.xlsx"
    df.to_excel(tracker_path, sheet_name="List", index=False)
    print(f"✅ Created tracker: {tracker_path}")
    return tracker_path


def create_text_file(vdr_dir):
    """Create a sample text file with known terms."""
    content = """This is a test document.

Acme Corporation is mentioned here.
Contact John Smith for more information.
Phone: (555) 123-4567

This file also contains Confidential information.

Some terms that should match via priority prefix:
- SpiritAirlines (should match 'spirit' prefix)
- AcmeCorp (should NOT match 'Acme Corporation' - different token)
"""
    path = vdr_dir / "sample_document.txt"
    with open(path, "w") as f:
        f.write(content)
    print(f"✅ Created text file: {path}")
    return path


def create_csv_file(vdr_dir):
    """Create a sample CSV file with known terms."""
    data = {
        "Name": ["John Smith", "Jane Doe", "Bob Wilson"],
        "Company": ["Acme Corporation", "Other Inc", "Test Corp"],
        "Phone": ["(555) 123-4567", "(555) 999-8888", "(555) 444-3333"],
        "Notes": ["Confidential data", "Public info", "General notes"],
    }
    df = pd.DataFrame(data)
    path = vdr_dir / "data.csv"
    df.to_csv(path, index=False)
    print(f"✅ Created CSV file: {path}")
    return path


def create_excel_file(vdr_dir):
    """Create a sample Excel file with known terms."""
    data = {
        "Contact": ["John Smith"],
        "Organization": ["Acme Corporation"],
        "SSN": ["123-45-6789"],
        "Status": ["Confidential"],
    }
    df = pd.DataFrame(data)
    path = vdr_dir / "spreadsheet.xlsx"
    df.to_excel(path, index=False)
    print(f"✅ Created Excel file: {path}")
    return path


def create_word_file(vdr_dir):
    """Create a sample Word document with known terms."""
    doc = Document()
    
    # Add content to body
    doc.add_heading("Test Document", 0)
    doc.add_paragraph("This document contains test data for Acme Corporation.")
    doc.add_paragraph("Contact: John Smith")
    doc.add_paragraph("This section is Confidential.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Phone"
    table.cell(1, 0).text = "John Smith"
    table.cell(1, 1).text = "(555) 123-4567"
    
    path = vdr_dir / "document.docx"
    doc.save(path)
    print(f"✅ Created Word file: {path}")
    return path


def create_powerpoint_file(vdr_dir):
    """Create a sample PowerPoint file with known terms."""
    prs = Presentation()
    
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Acme Corporation Presentation"
    subtitle.text = "By John Smith"
    
    # Content slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]
    title_shape.text = "Confidential Information"
    tf = body_shape.text_frame
    tf.text = "Contact: (555) 123-4567"
    
    path = vdr_dir / "presentation.pptx"
    prs.save(path)
    print(f"✅ Created PowerPoint file: {path}")
    return path


def create_priority_terms_file(fixtures_dir):
    """Create a file with priority terms for prefix matching."""
    terms = [
        "spirit",
        "acme",
        "test",
    ]
    path = fixtures_dir / "priority_terms.txt"
    with open(path, "w") as f:
        for term in terms:
            f.write(term + "\n")
    print(f"✅ Created priority terms file: {path}")
    return path


def create_all_fixtures():
    """Create all test fixtures."""
    print("Creating test fixtures...")
    print("=" * 60)
    
    fixtures_dir, vdr_dir = create_fixtures_dir()
    
    # Create tracker
    tracker_path = create_tracker_file(fixtures_dir)
    
    # Create VDR folder contents
    create_text_file(vdr_dir)
    create_csv_file(vdr_dir)
    create_excel_file(vdr_dir)
    create_word_file(vdr_dir)
    create_powerpoint_file(vdr_dir)
    
    # Create priority terms file
    priority_path = create_priority_terms_file(fixtures_dir)
    
    print("=" * 60)
    print(f"✅ All fixtures created in: {fixtures_dir}")
    print(f"   Tracker: {tracker_path}")
    print(f"   VDR folder: {vdr_dir}")
    print(f"   Priority terms: {priority_path}")
    
    return fixtures_dir, vdr_dir, tracker_path, priority_path


if __name__ == "__main__":
    create_all_fixtures()
